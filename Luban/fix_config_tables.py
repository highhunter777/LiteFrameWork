# -*- coding: utf-8 -*-
"""修正 2026-09-28 配置批的 xlsx：__tables__ 两行对齐（##标记 + 字符串true）+ 两张新数据表数据行重排。"""
import os
import openpyxl

ROOT = r"E:\unityProject\Test\Luban"
DATA = os.path.join(ROOT, "Data")

def set_row(ws, row_idx, values):
    for i, v in enumerate(values, start=1):
        ws.cell(row=row_idx, column=i, value=v)

# ---------- ① __tables__.xlsx：整行重写（A=## 标记，K=output，bool 用字符串 "true"） ----------
tp = os.path.join(DATA, "__tables__.xlsx")
wb = openpyxl.load_workbook(tp)
ws = wb.worksheets[0]
item_row = move_row = None
for row in ws.iter_rows():
    for c in row:
        if c.value == "TbItemConfig":
            item_row = c.row
        if c.value == "TbMovementConfig":
            move_row = c.row
assert item_row and move_row, f"行定位失败 item={item_row} move={move_row}"
# 列序：A ##var | B full_name | C value_type | D read_schema_from_file | E input | F index | G mode | H group | I comment | J tags | K output
set_row(ws, item_row, ["##", "TbItemConfig", "ItemConfig", "true", "#itemconfig.xlsx", "id", "", "c,s",
                       "道具表（类型 + 刷新/拾取/携带/使用 + 各类型效果数值）", "", "itemconfig"])
set_row(ws, move_row, ["##", "TbMovementConfig", "MovementConfig", "true", "#movementconfig.xlsx", "id", "", "c,s",
                       "移动数值表（单行：走跑冲/加减速/滑铲/空中/跳跃/钩爪/闪现；消费点 MovementConfig 静态类）", "", "movementconfig"])
wb.save(tp)
print("__tables__.xlsx rows realigned")

# ---------- ② #itemconfig.xlsx：数据行重排（type 为整型码，name 独立列） ----------
ic = openpyxl.Workbook()
s = ic.active
s.title = "Sheet"
set_row(s, 1, ["##var", "id", "type", "name", "spawn_interval", "max_alive", "pickup_radius", "carry_limit", "use_duration",
               "heal_amount", "shield_amount", "grenade_damage", "grenade_radius", "grenade_falloff",
               "emp_radius", "emp_duration", "radar_radius", "radar_duration", "projectile_speed", "travel_distance"])
set_row(s, 2, ["##type", "int", "int", "string", "float", "int", "float", "int", "float",
               "int", "int", "int", "float", "float",
               "float", "float", "float", "float", "float", "float"])
set_row(s, 3, ["##group", "", "c,s", "c,s", "c,s", "c,s", "c,s", "c,s", "c,s",
               "c,s", "c,s", "c,s", "c,s", "c,s",
               "c,s", "c,s", "c,s", "c,s", "c,s", "c,s"])
set_row(s, 4, ["##", "主键", "类型：1医疗包 2护盾电池 3手雷 4EMP 5雷达 6钩爪 7传送", "名称",
               "刷新间隔 s（0=不刷）", "场上最大存在数", "拾取半径 m", "携带上限", "使用耗时 s",
               "医疗包回复 HP", "护盾充能 mAh", "手雷伤害", "爆炸半径 m", "伤害衰减（0=无衰减）",
               "EMP 半径 m", "EMP 持续 s", "雷达半径 m", "雷达持续 s", "投掷物初速 m/s", "钩爪/传送距离 m"])
rows = [
    (1, 1, "医疗包",   15, 4, 1.5, 3, 1.5, 50, 0,  0,  0, 0,   0, 0, 0,  0, 0,  0),
    (2, 2, "护盾电池", 15, 4, 1.5, 3, 1.0, 0,  25, 0,  0, 0,   0, 0, 0,  0, 0,  0),
    (3, 3, "手雷",     20, 4, 1.5, 4, 0.6, 0,  0,  80, 5, 0.5, 0, 0, 0,  0, 20, 0),
    (4, 4, "EMP",      25, 2, 1.5, 2, 0.6, 0,  0,  0,  0, 0,   6, 3, 0,  0, 15, 0),
    (5, 5, "雷达",     30, 2, 1.5, 1, 0.8, 0,  0,  0,  0, 0,   0, 0, 40, 8, 0,  0),
    (6, 6, "钩爪",     40, 2, 1.5, 1, 0.5, 0,  0,  0,  0, 0,   0, 0, 0,  0, 0,  30),
    (7, 7, "传送器",   45, 2, 1.5, 1, 0.5, 0,  0,  0,  0, 0,   0, 0, 0,  0, 0,  20),
]
r = 5
for row in rows:
    set_row(s, r, [""] + list(row))
    r += 1
ic.save(os.path.join(DATA, "#itemconfig.xlsx"))
print("#itemconfig.xlsx data rows fixed (7 rows × 19 fields)")

# ---------- ③ #movementconfig.xlsx：数据行补 air_control（20 列对齐） ----------
mc = openpyxl.Workbook()
s = mc.active
s.title = "Sheet"
set_row(s, 1, ["##var", "id", "walk_speed", "run_speed", "sprint_speed", "acceleration", "sprint_duration",
               "slide_speed", "slide_friction", "slide_turn_penalty",
               "air_control", "gravity", "jump_speed", "double_jump_count", "double_jump_speed",
               "grapple_distance", "grapple_speed", "grapple_cooldown", "blink_distance", "blink_cooldown"])
set_row(s, 2, ["##type", "int", "float", "float", "float", "float", "float",
               "float", "float", "float",
               "float", "float", "float", "int", "float",
               "float", "float", "float", "float", "float"])
set_row(s, 3, ["##group", "", "c,s", "c,s", "c,s", "c,s", "c,s",
               "c,s", "c,s", "c,s",
               "c,s", "c,s", "c,s", "c,s", "c,s",
               "c,s", "c,s", "c,s", "c,s", "c,s"])
set_row(s, 4, ["##", "主键（单行表固定 1）", "走速 m/s", "跑速 m/s", "冲刺速度 m/s", "加速度 m/s²", "冲刺持续 s",
               "滑铲初速 m/s", "滑铲衰减 m/s²", "滑铲转向惩罚（0~1）",
               "空中控制系数（0~1）", "重力 m/s²（负值向下）", "跳跃初速 m/s", "二段跳次数", "二段跳速度 m/s",
               "钩爪距离 m", "钩爪速度 m/s", "钩爪冷却 s", "闪现距离 m", "闪现冷却 s"])
set_row(s, 5, ["", 1, 2.5, 5, 7.5, 25, 2.0,
               9, 8, 0.5,
               0.35, -20, 8, 1, 7,
               30, 20, 8, 10, 10])
mc.save(os.path.join(DATA, "#movementconfig.xlsx"))
print("#movementconfig.xlsx data row fixed (20 cols)")
print("DONE")
